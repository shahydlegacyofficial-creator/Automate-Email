"""Parsing, validating and de-duplicating recipient email addresses.

Recipients can come from pasted text (one per line, comma/semicolon separated,
``Name <email>`` style, or copied spreadsheet rows) or from CSV / TXT / XLSX files.
"""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from pathlib import Path

# Loose pattern used to *find* address-like tokens inside arbitrary text.
EMAIL_FIND_RE = re.compile(r"[A-Za-z0-9._%+'\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,}")

# Stricter pattern used to *validate* a single address.
EMAIL_VALID_RE = re.compile(
    r"^(?!\.)(?!.*\.\.)[A-Za-z0-9._%+'\-]+(?<!\.)"
    r"@(?:[A-Za-z0-9](?:[A-Za-z0-9\-]{0,61}[A-Za-z0-9])?\.)+[A-Za-z]{2,24}$"
)

# "John Smith <john@example.com>" or "<john@example.com>"
ANGLE_RE = re.compile(r'(?:"?([^"<>,;\n\t|]*?)"?\s*)?<\s*([^<>\s]+@[^<>\s]+)\s*>')

SPLIT_RE = re.compile(r"[,;\t|]")

EMAIL_HEADERS = ("email", "e-mail", "email address", "e-mail address", "mail", "email id")
NAME_HEADERS = ("name", "full name", "fullname", "customer name", "client name",
                "contact name", "customer", "client", "contact", "display name", "guest name")
FIRST_HEADERS = ("first name", "firstname", "first", "given name", "forename")
LAST_HEADERS = ("last name", "lastname", "last", "surname", "family name")

SUPPORTED_IMPORT_SUFFIXES = (".csv", ".txt", ".tsv", ".xlsx", ".xlsm")


def normalize_email(value: str) -> str:
    return (value or "").strip().strip("<>\"'").strip().rstrip(".").lower()


def is_valid_email(value: str) -> bool:
    return bool(EMAIL_VALID_RE.match((value or "").strip()))


def _clean_name(value: str) -> str:
    value = (value or "").strip().strip("\"'").strip()
    if not value or "@" in value:
        return ""
    if not re.search(r"[A-Za-zÀ-ɏЀ-ӿ؀-ۿ]", value):
        return ""  # pure numbers / punctuation are not names
    return re.sub(r"\s+", " ", value)


@dataclass
class Recipient:
    email: str
    name: str = ""

    @property
    def first_name(self) -> str:
        return self.name.split()[0] if self.name.strip() else ""

    @property
    def last_name(self) -> str:
        parts = self.name.split()
        return parts[-1] if len(parts) > 1 else ""

    def to_dict(self) -> dict:
        return {"email": self.email, "name": self.name}

    @classmethod
    def from_dict(cls, data: dict) -> "Recipient":
        return cls(email=normalize_email(data.get("email", "")), name=data.get("name", "") or "")


@dataclass
class ParseResult:
    recipients: list[Recipient] = field(default_factory=list)
    invalid: list[str] = field(default_factory=list)
    duplicates: int = 0

    def _add(self, email: str, name: str, seen: set[str]) -> None:
        email = normalize_email(email)
        if not email:
            return
        if not is_valid_email(email):
            if email not in self.invalid:
                self.invalid.append(email)
            return
        if email in seen:
            self.duplicates += 1
            # Fill in a name if an earlier occurrence had none.
            if name:
                for r in self.recipients:
                    if r.email == email and not r.name:
                        r.name = name
            return
        seen.add(email)
        self.recipients.append(Recipient(email=email, name=_clean_name(name)))


def _parse_line(line: str, result: ParseResult, seen: set[str]) -> None:
    line = line.strip()
    if not line or "@" not in line:
        return

    angle_matches = list(ANGLE_RE.finditer(line))
    if angle_matches:
        for m in angle_matches:
            result._add(m.group(2), m.group(1) or "", seen)
        # Anything outside the <...> blocks may hold more plain addresses.
        rest = ANGLE_RE.sub(" ", line)
        for email in EMAIL_FIND_RE.findall(rest):
            result._add(email, "", seen)
        return

    tokens = [t.strip() for t in SPLIT_RE.split(line) if t.strip()]
    emails_in_line = EMAIL_FIND_RE.findall(line)
    if len(emails_in_line) == 1:
        name = next((_clean_name(t) for t in tokens if _clean_name(t)), "")
        result._add(emails_in_line[0], name, seen)
    elif emails_in_line:
        for email in emails_in_line:
            result._add(email, "", seen)

    # Report tokens that look like addresses but are malformed (e.g. "john@gmail").
    for token in re.split(r"[\s,;\t|]+", line):
        token = token.strip("<>\"'()[]")
        if "@" in token and not EMAIL_FIND_RE.search(token):
            bad = normalize_email(token)
            if bad and bad not in result.invalid:
                result.invalid.append(bad)


def parse_text(text: str) -> ParseResult:
    result = ParseResult()
    seen: set[str] = set()
    for line in (text or "").splitlines():
        _parse_line(line, result, seen)
    return result


def _find_header(row: list[str], names: tuple[str, ...]) -> int:
    lowered = [(c or "").strip().lower() for c in row]
    for wanted in names:
        if wanted in lowered:
            return lowered.index(wanted)
    return -1


def parse_rows(rows: list[list[str]]) -> ParseResult:
    """Parse tabular data (CSV / Excel). Uses headers when present."""
    rows = [[("" if c is None else str(c)).strip() for c in row] for row in rows]
    rows = [r for r in rows if any(r)]
    result = ParseResult()
    seen: set[str] = set()
    if not rows:
        return result

    header = rows[0]
    email_idx = _find_header(header, EMAIL_HEADERS)
    if email_idx < 0:
        # Fuzzy: any header containing "mail"
        for i, cell in enumerate(header):
            if "mail" in cell.lower() and "@" not in cell:
                email_idx = i
                break

    if email_idx >= 0:
        name_idx = _find_header(header, NAME_HEADERS)
        first_idx = _find_header(header, FIRST_HEADERS)
        last_idx = _find_header(header, LAST_HEADERS)
        for row in rows[1:]:
            if email_idx >= len(row):
                continue
            cell = row[email_idx]
            found = EMAIL_FIND_RE.findall(cell) or ([cell] if "@" in cell else [])
            name = ""
            if 0 <= name_idx < len(row):
                name = row[name_idx]
            if not name:
                first = row[first_idx] if 0 <= first_idx < len(row) else ""
                last = row[last_idx] if 0 <= last_idx < len(row) else ""
                name = f"{first} {last}".strip()
            for email in found:
                result._add(email, name if len(found) == 1 else "", seen)
        return result

    for row in rows:
        _parse_line(",".join(row), result, seen)
    return result


def _read_text_file(path: Path) -> str:
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-16", "cp1252", "latin-1"):
        try:
            text = raw.decode(encoding)
            if encoding == "utf-16" and "\x00" in text:
                continue
            return text
        except (UnicodeDecodeError, UnicodeError):
            continue
    return raw.decode("utf-8", errors="ignore")


def parse_file(path: str | Path) -> ParseResult:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        combined = ParseResult()
        seen: set[str] = set()
        try:
            for sheet in wb.worksheets:
                rows = [list(r) for r in sheet.iter_rows(values_only=True)]
                sub = parse_rows(rows)
                combined.invalid.extend(i for i in sub.invalid if i not in combined.invalid)
                combined.duplicates += sub.duplicates
                for r in sub.recipients:
                    combined._add(r.email, r.name, seen)
        finally:
            wb.close()
        return combined

    text = _read_text_file(path)
    if suffix in (".csv", ".tsv"):
        sample = text[:4096]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel_tab if suffix == ".tsv" else csv.excel
        rows = list(csv.reader(io.StringIO(text), dialect))
        return parse_rows(rows)
    return parse_text(text)


def merge(existing: list[Recipient], incoming: list[Recipient]) -> tuple[list[Recipient], int, int]:
    """Merge *incoming* into *existing*. Returns (merged, added, duplicates)."""
    index = {r.email: r for r in existing}
    merged = list(existing)
    added = dupes = 0
    for r in incoming:
        current = index.get(r.email)
        if current is not None:
            dupes += 1
            if r.name and not current.name:
                current.name = r.name
            continue
        copy = Recipient(r.email, r.name)
        index[r.email] = copy
        merged.append(copy)
        added += 1
    return merged, added, dupes


def write_csv(path: str | Path, recipients: list[Recipient]) -> None:
    with open(path, "w", newline="", encoding="utf-8-sig") as fh:
        writer = csv.writer(fh)
        writer.writerow(["Email", "Name"])
        for r in recipients:
            writer.writerow([r.email, r.name])


def parse_suppression(text: str) -> set[str]:
    """Do-not-email list: addresses and/or whole domains (``@domain.com``)."""
    entries: set[str] = set()
    for token in re.split(r"[\s,;]+", text or ""):
        token = token.strip().lower()
        if token:
            entries.add(token)
    return entries


def is_suppressed(email: str, entries: set[str]) -> bool:
    email = email.lower()
    if email in entries:
        return True
    domain = email.rsplit("@", 1)[-1]
    return f"@{domain}" in entries or domain in entries
