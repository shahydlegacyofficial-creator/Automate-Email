"""Collect customer data and past Travel Essence emails from Outlook."""

from __future__ import annotations

import csv
from datetime import datetime, time as dtime

from PySide6.QtCore import QDate, Qt
from PySide6.QtWidgets import (
    QAbstractItemView, QComboBox, QDateEdit, QFileDialog, QGridLayout, QHBoxLayout, QHeaderView,
    QLineEdit, QMessageBox, QProgressBar, QSpinBox, QTabWidget, QVBoxLayout, QWidget,
)

from ...core import outlook_data
from ...core.outlook_data import CollectOptions, CollectResult, FOLDER_CHOICES
from ...core.recipients import Recipient, is_valid_email
from ...core.senders import outlook_available
from ...core import graph
from ..microsoft import MicrosoftAccountCard
from ..widgets import Banner, Card, Page, Pill, button, field_label, label, make_table, table_item
from ..workers import run_task


def _fmt(dt: datetime | None) -> str:
    return dt.strftime("%Y-%m-%d %H:%M") if dt else ""


class OutlookPage(Page):
    def __init__(self, state):
        super().__init__("Outlook data",
                         "Collect customer addresses and the Travel Essence emails you have sent - from the "
                         "Outlook app on this PC or online from Microsoft 365 / the new Outlook.")
        self.state = state
        self.result = CollectResult()
        self._task = None
        self._custom: tuple[str, str, str] | None = None

        ok, reason = outlook_available()
        self.banner = Banner(
            f"<b>The classic Outlook app is not available.</b> {reason} If you use the new Outlook, choose "
            "<b>Microsoft 365 / new Outlook (online)</b> under Connect to.")
        self.banner.setVisible(False)
        self._classic_ok = ok
        self.content.addWidget(self.banner)
        self.ms_card = MicrosoftAccountCard(state, "Microsoft 365 / new Outlook connection")
        self.content.addWidget(self.ms_card)

        # ---- source card
        source = Card("Source", "Choose where to read from. Your own addresses are always excluded.")
        grid = QGridLayout()
        grid.setHorizontalSpacing(14)
        grid.setVerticalSpacing(6)

        self.source_box = QComboBox()
        self.source_box.addItem("Outlook app on this PC (classic)", "classic")
        self.source_box.addItem("Microsoft 365 / new Outlook (online)", "graph")
        self.source_box.setCurrentIndex(max(0, self.source_box.findData(state.settings.get("outlook_source", "classic"))))
        self.source_box.currentIndexChanged.connect(self._source_changed)

        self.store_box = QComboBox()
        self.store_box.addItem("All accounts", "")
        self.store_box.setMinimumWidth(220)
        connect_btn = button("Connect", "refresh")
        connect_btn.setToolTip("Load the accounts configured in Outlook")
        connect_btn.clicked.connect(self.connect_outlook)
        self.store_widget = QWidget()
        store_row = QHBoxLayout(self.store_widget)
        store_row.setContentsMargins(0, 0, 0, 0)
        store_row.addWidget(self.store_box, 1)
        store_row.addWidget(connect_btn)
        self.store_caption = field_label("ACCOUNT / MAILBOX")

        self.folder_box = QComboBox()
        for key, text in FOLDER_CHOICES:
            self.folder_box.addItem(text, key)
        self.folder_box.currentIndexChanged.connect(self._folder_changed)

        self.date_from = QDateEdit(QDate.currentDate().addMonths(-6))
        self.date_from.setCalendarPopup(True)
        self.date_from.setDisplayFormat("dd MMM yyyy")
        self.date_to = QDateEdit(QDate.currentDate())
        self.date_to.setCalendarPopup(True)
        self.date_to.setDisplayFormat("dd MMM yyyy")

        self.subject_filter = QLineEdit()
        self.subject_filter.setPlaceholderText("e.g. Travel Essence, Booking, Quotation")
        self.exclude = QLineEdit(state.settings.get("outlook_exclude", ""))
        self.exclude.setPlaceholderText("@travelessence.com, noreply@…")
        self.max_items = QSpinBox()
        self.max_items.setRange(50, 100000)
        self.max_items.setSingleStep(500)
        self.max_items.setValue(5000)

        fields = [("CONNECT TO", self.source_box), (self.store_caption, self.store_widget),
                  ("FOLDER", self.folder_box), ("FROM DATE", self.date_from), ("TO DATE", self.date_to),
                  ("SUBJECT CONTAINS", self.subject_filter), ("EXCLUDE ADDRESSES / DOMAINS", self.exclude),
                  ("MAX EMAILS", self.max_items)]
        for i, (caption, widget) in enumerate(fields):
            r, c = divmod(i, 4)
            grid.addWidget(caption if isinstance(caption, QWidget) else field_label(caption), r * 2, c)
            grid.addWidget(widget, r * 2 + 1, c)
        for c in range(4):
            grid.setColumnStretch(c, 1)
        source.add(grid)

        run_row = QHBoxLayout()
        self.collect_btn = button("Collect data", "database", "primary")
        self.collect_btn.clicked.connect(self.collect)
        self.cancel_btn = button("Cancel", "stop")
        self.cancel_btn.clicked.connect(self.cancel)
        self.cancel_btn.setEnabled(False)
        self.progress = QProgressBar()
        self.progress.setObjectName("ThinProgress")
        self.progress.setTextVisible(False)
        self.progress.setRange(0, 1)
        self.progress.setValue(0)
        self.progress_label = label("", "Hint")
        run_row.addWidget(self.collect_btn)
        run_row.addWidget(self.cancel_btn)
        prog_col = QVBoxLayout()
        prog_col.setSpacing(4)
        prog_col.addWidget(self.progress_label)
        prog_col.addWidget(self.progress)
        run_row.addLayout(prog_col, 1)
        source.add(run_row)
        self.content.addWidget(source)

        # ---- results
        results = Card("Results")
        self.summary_pill = Pill("Nothing collected yet")
        results.add_action(self.summary_pill)
        self.tabs = QTabWidget()

        # contacts tab
        contacts_tab = QWidget()
        ct = QVBoxLayout(contacts_tab)
        ct.setContentsMargins(0, 10, 0, 0)
        ct.setSpacing(10)
        ctools = QHBoxLayout()
        self.contact_search = QLineEdit()
        self.contact_search.setPlaceholderText("Search contacts…")
        self.contact_search.setClearButtonEnabled(True)
        self.contact_search.textChanged.connect(lambda t: self._filter(self.contacts_table, t))
        add_sel = button("Add selected to recipients", "user-plus", "primary")
        add_sel.clicked.connect(lambda: self.add_to_recipients(selected_only=True))
        add_all = button("Add all", "users")
        add_all.clicked.connect(lambda: self.add_to_recipients(selected_only=False))
        exp_c = button("Export", "download")
        exp_c.clicked.connect(self.export_contacts)
        ctools.addWidget(self.contact_search, 1)
        for b in (add_sel, add_all, exp_c):
            ctools.addWidget(b)
        ct.addLayout(ctools)
        self.contacts_table = make_table(["Email", "Name", "Company", "Phone", "Emails", "Last contact",
                                          "Last subject"], 6)
        self.contacts_table.setSelectionMode(QAbstractItemView.SelectionMode.ExtendedSelection)
        self.contacts_table.setColumnWidth(0, 240)
        self.contacts_table.setColumnWidth(1, 170)
        self.contacts_table.setColumnWidth(2, 130)
        self.contacts_table.setColumnWidth(3, 110)
        self.contacts_table.setColumnWidth(4, 70)
        self.contacts_table.setColumnWidth(5, 130)
        self.contacts_table.setMinimumHeight(380)
        self.contacts_table.setSortingEnabled(True)
        ct.addWidget(self.contacts_table, 1)
        self.tabs.addTab(contacts_tab, "Customers")

        # messages tab
        messages_tab = QWidget()
        mt = QVBoxLayout(messages_tab)
        mt.setContentsMargins(0, 10, 0, 0)
        mt.setSpacing(10)
        mtools = QHBoxLayout()
        self.message_search = QLineEdit()
        self.message_search.setPlaceholderText("Search emails…")
        self.message_search.setClearButtonEnabled(True)
        self.message_search.textChanged.connect(lambda t: self._filter(self.messages_table, t))
        use_tpl = button("Use as template", "compose", "primary")
        use_tpl.setToolTip("Load the selected email's subject, content and attachments into Compose")
        use_tpl.clicked.connect(self.use_as_template)
        exp_m = button("Export", "download")
        exp_m.clicked.connect(self.export_messages)
        mtools.addWidget(self.message_search, 1)
        mtools.addWidget(use_tpl)
        mtools.addWidget(exp_m)
        mt.addLayout(mtools)
        self.messages_table = make_table(["Date", "Folder", "From", "To", "Subject", "Attachments"], 4)
        self.messages_table.setColumnWidth(0, 130)
        self.messages_table.setColumnWidth(1, 100)
        self.messages_table.setColumnWidth(2, 200)
        self.messages_table.setColumnWidth(3, 240)
        self.messages_table.horizontalHeader().setSectionResizeMode(5, QHeaderView.ResizeMode.ResizeToContents)
        self.messages_table.setMinimumHeight(380)
        self.messages_table.setSortingEnabled(True)
        self.messages_table.doubleClicked.connect(lambda _: self.use_as_template())
        mt.addWidget(self.messages_table, 1)
        self.tabs.addTab(messages_tab, "Emails")

        results.add(self.tabs, 1)
        self.content.addWidget(results, 1)
        self._source_changed()

    # ------------------------------------------------------------ outlook
    def _online(self) -> bool:
        return self.source_box.currentData() == "graph"

    def _source_changed(self, *_):
        online = self._online()
        self.ms_card.setVisible(online)
        self.banner.setVisible(not online and not self._classic_ok)
        self.store_caption.setVisible(not online)
        self.store_widget.setVisible(not online)
        custom = self.folder_box.findData("custom")
        item = self.folder_box.model().item(custom)
        if item is not None:
            item.setEnabled(not online)
        if online and self.folder_box.currentData() == "custom":
            self.folder_box.setCurrentIndex(0)
        if self.state.settings.get("outlook_source") != self.source_box.currentData():
            self.state.save_settings({"outlook_source": self.source_box.currentData()})

    def connect_outlook(self):
        def done(info):
            current = self.store_box.currentData()
            self.store_box.clear()
            self.store_box.addItem("All accounts", "")
            for store in info.get("stores", []):
                self.store_box.addItem(store, store)
            idx = self.store_box.findData(current)
            self.store_box.setCurrentIndex(max(0, idx))
            names = ", ".join(a["smtp"] or a["name"] for a in info.get("accounts", []))
            self.progress_label.setText(f"Connected · accounts: {names or 'none found'}")
            self.state.notify("Connected to Outlook", "success")

        self.progress_label.setText("Connecting to Outlook…")
        self._task = run_task(lambda _ctx: outlook_data.list_accounts(), done, self._error)

    def _folder_changed(self):
        if self.folder_box.currentData() != "custom" or self._online():
            return

        def done(picked):
            if picked is None:
                self.folder_box.setCurrentIndex(0)
                return
            self._custom = picked
            label_text = f"Folder: {picked[2]}"
            idx = self.folder_box.findData("custom")
            self.folder_box.setItemText(idx, label_text)

        self._task = run_task(lambda _ctx: outlook_data.pick_folder(), done, self._error)

    def _options(self) -> CollectOptions:
        folder = self.folder_box.currentData()
        options = CollectOptions(
            store=self.store_box.currentData() or "",
            folder=folder,
            date_from=datetime.combine(self.date_from.date().toPython(), dtime.min),
            date_to=datetime.combine(self.date_to.date().toPython(), dtime.max),
            subject_contains=self.subject_filter.text(),
            exclude=[e.strip().lower() for e in self.exclude.text().replace(";", ",").split(",") if e.strip()],
            max_items=self.max_items.value(),
        )
        if folder == "custom" and self._custom:
            options.custom_entry_id, options.custom_store_id = self._custom[0], self._custom[1]
        return options

    def collect(self):
        if self._task and self._task.is_running():
            return
        options = self._options()
        if options.folder == "custom" and not options.custom_entry_id:
            QMessageBox.information(self, "Choose folder", "Pick a folder first.")
            return
        if self.exclude.text() != self.state.settings.get("outlook_exclude", ""):
            self.state.save_settings({"outlook_exclude": self.exclude.text()})

        if self._online():
            if not self.ms_card.is_signed_in():
                QMessageBox.information(self, "Sign in first", "Sign in to Microsoft 365 / new Outlook above "
                                                               "to read your mailbox online.")
                return
            client = self.ms_card.client()
            work = lambda ctx: graph.collect(client, options, ctx.progress, ctx.cancelled)  # noqa: E731
        else:
            work = lambda ctx: outlook_data.collect(options, ctx.progress, ctx.cancelled)  # noqa: E731

        self.collect_btn.setEnabled(False)
        self.cancel_btn.setEnabled(True)
        self.progress.setRange(0, 0)
        self.progress_label.setText("Reading your mailbox… this can take a minute for large mailboxes.")

        def progress(current, total, message):
            if total > 0:
                self.progress.setRange(0, total)
                self.progress.setValue(min(current, total))
            self.progress_label.setText(message)

        self._task = run_task(work, self._collected, self._error, progress)

    def cancel(self):
        if self._task:
            self._task.cancel()
            self.progress_label.setText("Cancelling…")

    def _error(self, message: str):
        self.collect_btn.setEnabled(True)
        self.cancel_btn.setEnabled(False)
        self.progress.setRange(0, 1)
        self.progress.setValue(0)
        self.progress_label.setText("")
        if self.folder_box.currentData() == "custom" and not self._custom:
            self.folder_box.setCurrentIndex(0)
        QMessageBox.critical(self, "Outlook", message)

    def _collected(self, result: CollectResult):
        self.result = result
        self.collect_btn.setEnabled(True)
        self.cancel_btn.setEnabled(False)
        self.progress.setRange(0, 1)
        self.progress.setValue(1)
        self.progress_label.setText(f"Scanned {result.scanned:,} items")
        self.summary_pill.set(f"{len(result.contacts):,} customers · {len(result.messages):,} emails", "accent")
        self.tabs.setTabText(0, f"Customers ({len(result.contacts):,})")
        self.tabs.setTabText(1, f"Emails ({len(result.messages):,})")

        t = self.contacts_table
        t.setSortingEnabled(False)
        t.setRowCount(len(result.contacts))
        for row, c in enumerate(result.contacts):
            t.setItem(row, 0, table_item(c.email, data=c.email))
            t.setItem(row, 1, table_item(c.name))
            t.setItem(row, 2, table_item(c.company))
            t.setItem(row, 3, table_item(c.phone))
            count = table_item("")
            count.setData(Qt.ItemDataRole.DisplayRole, c.interactions)
            t.setItem(row, 4, count)
            t.setItem(row, 5, table_item(_fmt(c.last_contact)))
            t.setItem(row, 6, table_item(c.last_subject))
        t.setSortingEnabled(True)

        m = self.messages_table
        m.setSortingEnabled(False)
        m.setRowCount(len(result.messages))
        for row, msg in enumerate(result.messages):
            m.setItem(row, 0, table_item(_fmt(msg.date), data=(msg.entry_id, msg.store_id)))
            m.setItem(row, 1, table_item(msg.folder))
            m.setItem(row, 2, table_item(msg.sender))
            to_item = table_item(msg.recipients)
            to_item.setToolTip(msg.recipients)
            m.setItem(row, 3, to_item)
            m.setItem(row, 4, table_item(msg.subject))
            m.setItem(row, 5, table_item(msg.attachments or ""))
        m.setSortingEnabled(True)
        self.state.notify(f"Collected {len(result.contacts):,} customers and {len(result.messages):,} emails",
                          "success")

    # ------------------------------------------------------------ actions
    @staticmethod
    def _filter(table, text: str):
        needle = text.strip().lower()
        for row in range(table.rowCount()):
            hay = " ".join((table.item(row, c).text() if table.item(row, c) else "")
                           for c in range(table.columnCount())).lower()
            table.setRowHidden(row, bool(needle) and needle not in hay)

    def add_to_recipients(self, selected_only: bool):
        t = self.contacts_table
        if selected_only:
            rows = sorted({i.row() for i in t.selectedItems()})
            if not rows:
                QMessageBox.information(self, "Add recipients", "Select one or more customers first "
                                                                "(Ctrl+A selects all).")
                return
        else:
            rows = [r for r in range(t.rowCount()) if not t.isRowHidden(r)]
        recipients = []
        for r in rows:
            email = t.item(r, 0).text()
            if is_valid_email(email):
                recipients.append(Recipient(email, t.item(r, 1).text()))
        if not recipients:
            return
        added, dupes = self.state.add_recipients(recipients)
        self.state.notify(f"{added:,} added to recipients" + (f" · {dupes} already in the list" if dupes else ""),
                          "success")

    def use_as_template(self):
        rows = sorted({i.row() for i in self.messages_table.selectedItems()})
        if not rows:
            QMessageBox.information(self, "Use as template", "Select an email in the list first.")
            return
        entry_id, store_id = self.messages_table.item(rows[0], 0).data(Qt.ItemDataRole.UserRole)
        if QMessageBox.question(self, "Use as template",
                                "Replace the email in Compose with the selected Outlook email?") != \
                QMessageBox.StandardButton.Yes:
            return

        def done(data):
            self.state.replace_draft(data["subject"], data["html"], data["attachments"])
            note = ""
            if data.get("skipped"):
                note = f" ({len(data['skipped'])} attachment(s) skipped - only JPG, PNG and PDF are supported)"
            self.state.notify("Email loaded into Compose" + note, "success")
            self.state.navigate.emit("compose")

        self.progress_label.setText("Loading email from Outlook…")
        if store_id == "graph":
            client = self.ms_card.client()
            loader = lambda _ctx: graph.load_message(client, entry_id)  # noqa: E731
        else:
            loader = lambda _ctx: outlook_data.load_message(entry_id, store_id)  # noqa: E731
        self._task = run_task(loader, done, self._error)

    def _export(self, table, default_name: str):
        if table.rowCount() == 0:
            QMessageBox.information(self, "Export", "Nothing to export yet.")
            return
        path, chosen = QFileDialog.getSaveFileName(self, "Export", default_name,
                                                   "Excel workbook (*.xlsx);;CSV (*.csv)")
        if not path:
            return
        headers = [table.horizontalHeaderItem(c).text() for c in range(table.columnCount())]
        rows = [[table.item(r, c).text() if table.item(r, c) else "" for c in range(table.columnCount())]
                for r in range(table.rowCount()) if not table.isRowHidden(r)]
        try:
            if path.lower().endswith(".xlsx") or ("xlsx" in chosen and not path.lower().endswith(".csv")):
                if not path.lower().endswith(".xlsx"):
                    path += ".xlsx"
                from openpyxl import Workbook
                from openpyxl.styles import Font

                wb = Workbook()
                ws = wb.active
                ws.append(headers)
                for cell in ws[1]:
                    cell.font = Font(bold=True)
                for row in rows:
                    ws.append(row)
                for col, header in enumerate(headers, start=1):
                    width = max([len(header)] + [len(r[col - 1]) for r in rows[:500]]) + 2
                    ws.column_dimensions[ws.cell(1, col).column_letter].width = min(width, 60)
                wb.save(path)
            else:
                with open(path, "w", newline="", encoding="utf-8-sig") as fh:
                    writer = csv.writer(fh)
                    writer.writerow(headers)
                    writer.writerows(rows)
        except OSError as exc:
            QMessageBox.critical(self, "Export failed", str(exc))
            return
        self.state.notify("Exported successfully", "success")

    def export_contacts(self):
        self._export(self.contacts_table, "travel-essence-customers.xlsx")

    def export_messages(self):
        self._export(self.messages_table, "travel-essence-emails.xlsx")
